"""Table source: find the row with the quote, draw it and its neighbours as a small table PDF."""

# Where the Excel format emulation stops: the codes handled are what is here now (round half up,
# zero padding, quoted words, the default date in this computer's order, percent, and the common
# number and date codes). Any other code shows the stored value, and that is the answer, not a bug.

import csv
import datetime
import html
import locale
import re
from decimal import ROUND_HALF_UP, Decimal

import openpyxl
import pymupdf
import xlrd

from claim import cache

AROUND = 4  # rows shown above and below the hit
RAW = "\x1f"  # between what Excel shows and the stored value in one cell; split() counts it as a space
KINDS = (".csv", ".tsv", ".xlsx", ".xlsm", ".xls")
BOOKS = (".xlsx", ".xlsm", ".xls")  # the kinds with more than one sheet
DATE_CODES = r"min|am/pm|yyyy|yy|mmmmm|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss"


def pick(path, quote, sheet):
    """All the rows, which of them to show (the header, then the hit row and its neighbours), and the hit row."""
    rows = read_rows(path, sheet)
    if not rows:
        raise ValueError("the table is empty")
    hit = max(range(len(rows)), key=lambda i: score(rows[i], quote))
    head = header(rows, hit)
    keep = range(max(hit - AROUND, head + 1), min(hit + AROUND + 1, len(rows)))
    return rows, [head] + [i for i in keep if i != head], hit  # always show the header row first


def header(rows, hit):
    """Which row names the columns: the first one above the hit with two or more filled cells, so a
    title on its own line ("Q3 travel expenses") and blank rows under it are passed over. Else row 1."""
    return next((i for i in range(hit) if sum(1 for v in rows[i] if v.strip()) >= 2), 0)


def score(row, quote):
    """How well a row fits the quote: holding all of it beats any number of shared words."""
    text = squash(" ".join(row))
    if within(quote, text):
        return 10 ** 6
    return len(set(squash(quote).split()) & set(text.split()))


def within(needle, text):
    """Whether the text holds the needle, spacing and case aside. A number must stand on its
    own, sign included: "2.5" is not inside "12.5%" or "-2.5", and "2100" is not inside "21000"."""
    needle, text = squash(needle), squash(text)
    if not needle:
        return False
    before = r"(?<!\d)(?<!\d[.,])(?<!\s-)(?<!^-)" if needle[0].isdigit() else ""
    after = r"(?!\d|[.,]\d)" if needle[-1].isdigit() else ""
    return re.search(before + re.escape(needle) + after, text) is not None


def pick_sheet(path, quote):
    """The name of the sheet whose best row fits the quote best; the first sheet on a tie."""
    return max(read_sheets(path), key=lambda s: max((score(row, quote) for row in s[1]), default=0))[0]


_last = {}  # the file read last, as {(path, changed, size): sheets}: one claim reads a table several times


def read_sheets(path):
    """Every sheet of the file as (name, rows); a CSV is one unnamed sheet. Read once per file version,
    and kept on disk (cache.py), so the next claim or scan on the same file doesn't open it again."""
    stamp = (str(path), path.stat().st_mtime_ns, path.stat().st_size)
    if stamp not in _last:
        _last.clear()
        _last[stamp] = cache.kept(path, "sheets", lambda: sheets(path) if path.suffix.lower() in BOOKS else [("", csv_rows(path))])
    return _last[stamp]


def sheets(path):
    """Every sheet of a workbook as (name, rows). .xls is the pre-2007 format, which needs xlrd."""
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path)  # a date comes as a day count (45297), so it is turned into a date first
        value = lambda c: xlrd.xldate_as_datetime(c.value, book.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value
        return [(ws.name, [[cell(value(c)) for c in ws.row(i)] for i in range(ws.nrows)]) for ws in book.sheets()]
    # A formula cell holds its last result only if Excel saved the file; a workbook written by a script
    # has none. So the file is opened twice, values and formulas, and an empty result shows the formula.
    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    return [(ws.title, [[shown(c.value if c.value is not None else f.value, c.number_format)
                         for c, f in zip(row, frow)] for row, frow in zip(ws.iter_rows(), fs.iter_rows())])
            for ws, fs in zip(values.worksheets, formulas.worksheets)]


def cell(v):
    """A cell as text. xlrd gives every number as a float, so 1200.0 is written 1200;
    a date with no time of day is written without the 00:00:00."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).removesuffix(" 00:00:00")


def shown(v, fmt):
    """A cell as Excel shows it ("$1,200.00", "5-Jan-24") for the common number and date
    formats. When that reads differently from the stored value, the stored one follows
    after RAW, so a quote in either form is found and the picture shows both."""
    raw = cell(v)
    try:
        display = as_date(v, fmt) if isinstance(v, datetime.date) else as_number(v, fmt)
    except (ValueError, IndexError):
        display = None
    return raw if not display or display == raw else f"{display}{RAW}{raw}"


def as_number(v, fmt):
    """The common codes only: a currency sign, thousands commas, decimal places, zero padding,
    a percent sign, brackets for a negative, quoted text. Anything else (exponents, fractions,
    "?" spaces) gives None, so the picture never shows a value Excel wouldn't."""
    if not fmt or fmt == "General" or not isinstance(v, (int, float)) or isinstance(v, bool):
        return None
    sections = fmt.split(";")  # positive; negative; zero
    section = sections[2] if v == 0 and len(sections) > 2 else sections[1] if v < 0 and len(sections) > 1 else sections[0]
    section = re.sub(r"\[\$([^-\]]*)[^\]]*\]", r"\1", section)       # [$€-407] is the currency €
    section = re.sub(r"\[[^\]]*\]|_.|\*.", "", section).replace("\\", "")  # colours, spacing
    section, literals = protect(section)                                # "kg" is the text kg, kept whole
    digits = re.search(r"[#0][#0,]*(\.0+)?", section)
    if not digits or re.search(r"[#0?/E@]", section[:digits.start()] + section[digits.end():]):
        return None
    places = len(digits[1] or ".") - 1
    thousands = len(digits[0]) - len(digits[0].rstrip(","))  # #,##0, shows 1,000,000 as 1,000
    n = Decimal(str(abs(v) * 100 if "%" in section else abs(v))) / 1000 ** thousands
    n = n.quantize(Decimal(1).scaleb(-places), ROUND_HALF_UP)
    text = f"{n:,.{places}f}" if "," in digits[0] else f"{n:.{places}f}"
    pad = digits[0].partition(".")[0].count("0")  # 0000 pads to four figures
    if pad > 1 and "," not in digits[0]:
        whole, dot, frac = text.partition(".")
        text = whole.zfill(pad) + dot + frac
    sign = "-" if v < 0 and len(sections) == 1 else ""
    return restore(sign + section[:digits.start()] + text + section[digits.end():], literals)


def as_date(v, fmt):
    if not fmt or fmt == "General":
        return None
    if fmt == "mm-dd-yy":  # Excel's default date: shown the way this computer writes a short date
        fmt = short_date()
    fmt, literals = protect(re.sub(r"\[[^\]]*\]|\\", "", fmt.split(";")[0]))
    fmt = re.sub(r"(?<=[h:])mm|mm(?=:s)", "min", fmt, flags=re.IGNORECASE)  # mm next to hours or seconds is minutes
    if re.search(r"[a-z]", re.sub(DATE_CODES, "", fmt, flags=re.IGNORECASE), re.IGNORECASE):
        return None  # a code not handled here (an era, A/P): the stored value, never a guess
    hour, minute, second = getattr(v, "hour", 0), getattr(v, "minute", 0), getattr(v, "second", 0)
    if "am/pm" in fmt.lower():
        hour = hour % 12 or 12
    parts = {"min": f"{minute:02d}", "am/pm": "AM" if getattr(v, "hour", 0) < 12 else "PM",
             "yyyy": f"{v.year}", "yy": f"{v.year % 100:02d}", "mmmmm": v.strftime("%b")[0], "mmmm": v.strftime("%B"),
             "mmm": v.strftime("%b"), "mm": f"{v.month:02d}", "m": str(v.month), "dddd": v.strftime("%A"),
             "ddd": v.strftime("%a"), "dd": f"{v.day:02d}", "d": str(v.day), "hh": f"{hour:02d}", "h": str(hour),
             "ss": f"{second:02d}"}
    shown_ = re.sub(DATE_CODES, lambda m: parts[m[0].lower()], fmt, flags=re.IGNORECASE)
    return restore(shown_, literals)


def short_date():
    """Excel's default date code means "the computer's short date": 4/12/2023 on a US machine,
    12/04/2023 in the UK. The computer's own setting says which of the day and month comes first."""
    try:
        locale.setlocale(locale.LC_TIME, "")
        order = locale.nl_langinfo(locale.D_FMT)
    except (locale.Error, AttributeError):
        order = "%m/%d/%Y"
    return "m/d/yyyy" if order.find("%m") < order.find("%d") else "d/m/yyyy"


def protect(fmt):
    """Quoted text in a format code taken out (a placeholder stands in), so its letters aren't read as codes."""
    literals = re.findall(r'"([^"]*)"', fmt)
    return re.sub(r'"[^"]*"', "\x00", fmt), literals


def restore(text, literals):
    for literal in literals:
        text = text.replace("\x00", literal, 1)
    return text


def squash(text):
    return " ".join(text.lower().split())


def load(path, quote, sheet=None):
    rows, shown, _ = pick(path, quote, sheet)

    width = max((c + 1 for i in shown for c, v in enumerate(rows[i]) if v), default=0)  # empty columns at the right are left out
    cells = "".join(
        f"<tr><td class='num'>{i + 1}</td>" + "".join(f"<td>{td(c)}</td>" for c in rows[i][:width]) + "</tr>"
        for i in shown)
    table = f"<table>{cells}</table>"
    css = ("* {font-family: sans-serif; font-size: 10pt} table {border-collapse: collapse}"
           " td {border: 1px solid #ccc; padding: 3px 8px; white-space: nowrap} td.num {color: #888; background: #f3f3f3}"
           " tr:first-child td {font-weight: bold}")

    # Lay the table out once on a big scratch page to see how much room it takes,
    # then again on a page that size, so the picture has no empty space to the right.
    scratch = pymupdf.open().new_page(width=2000, height=2000)
    scratch.insert_htmlbox(scratch.rect + (30, 30, -30, -30), table, css=css)
    used = pymupdf.Rect()
    for block in scratch.get_text("blocks"):
        used |= block[:4]
    doc = pymupdf.open()
    page = doc.new_page(width=used.x1 + 80, height=used.y1 + 40)
    page.insert_htmlbox(pymupdf.Rect(30, 30, used.x1 + 50, used.y1 + 10), table, css=css)
    return doc


def td(text):
    """A cell's HTML: what Excel shows, and the stored value small and grey under it when they differ."""
    shown_, _, raw = text.partition(RAW)
    return html.escape(shown_) + (f"<br><span style='color:#999;font-size:8pt'>{html.escape(raw)}</span>" if raw else "")


def excerpt(path, quote, sheet, matched):
    """The rows the picture shows, as plain cells, and which one holds the matched text (None if none)."""
    rows, shown, _ = pick(path, quote, sheet)
    body = shown[1:]
    hit = next((k for k, i in enumerate(body) if matched and holds(rows[i], matched)), None)
    # The cells of the hit row the match covers: one holding all of it, or each a piece of a match across cells.
    row = rows[body[hit]] if hit is not None else []
    cells = [c for c, v in enumerate(row) if any(squash(p) and (within(matched, p) or f" {squash(p)} " in f" {squash(matched)} ")
                                                 for p in [v, *v.split(RAW)])]
    plain = [[v.partition(RAW)[0] for v in rows[i]] for i in body]  # the card shows what Excel shows
    return {"header": [v.partition(RAW)[0] for v in rows[shown[0]]], "rows": plain, "numbers": [i + 1 for i in body], "hit": hit, "cells": cells}


def holds(row, matched):
    """Whether a row holds the matched text, or every value of a row match ("Malmo · 693968")."""
    return all(within(part, " ".join(row)) for part in matched.split(" · "))


def row_values(path, quote, sheet):
    """For a quote made of values from one row that don't sit side by side ("INV07021335 226548"):
    the row's number and the cell values in it, when together they hold every word of the quote.
    None when any word of the quote is in none of them, so a made-up value never passes."""
    rows, _, hit = pick(path, quote, sheet)
    words = [w.strip(",;:()") for w in squash(quote).split()]
    said = f" {' '.join(words)} "
    parts = list(dict.fromkeys(p for v in rows[hit] for p in v.split(RAW) if squash(p) and f" {squash(p)} " in said))
    covered = {w for p in parts for w in squash(p).split()}
    if len(parts) < 2 or not all(w in covered for w in words if w):
        return None
    return hit + 1, parts


def row_rects(page, number, values):
    """Where the values sit on the drawn table, on the row numbered `number` only. Row numbers
    are the first column, so a row runs from its number down to the next row's number.
    None if any value isn't on that row."""
    words = page.get_text("words")
    left = min(w[0] for w in words)
    firsts = [w for w in words if w[0] < left + 1]  # the row-number column
    top = next(w[1] for w in firsts if w[4] == str(number))
    bottom = min([w[1] for w in firsts if w[1] > top + 1], default=page.rect.y1)
    on_row = [w for w in words if top - 2 <= w[1] < bottom and w[0] > left + 1]  # never the row number itself
    # A one-word value must be a whole word of the row ("2" is not the "2" in "2100"); longer values are searched for.
    found = [[pymupdf.Rect(w[:4]) for w in on_row if w[4].strip(".,;:()").lower() == v.lower()] if " " not in v
             else [r for r in page.search_for(v) if top - 2 <= r.y0 < bottom and r.x0 > left + 1] for v in values]
    return None if not all(found) else [r for rects in found for r in rects]


def read_rows(path, sheet):
    found = read_sheets(path)
    return dict(found)[sheet] if sheet else found[0][1]


def csv_rows(path):
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [row for row in csv.reader(f, delimiter=delimiter)]
